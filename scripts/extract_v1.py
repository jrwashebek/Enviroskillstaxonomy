#!/usr/bin/env python3
"""Extract the EMST v1 baseline database into tidy, reproducible tables.

Reads the v1 baseline workbook (copyB of 20260223_SkillsTaxonomy_Core_V3_reconciled, see
docs/audit/data-audit.md section 1 for why) plus the three auxiliary exposure workbooks that hold
per-task rationale, and writes CSV tables + one combined JSON to data/v1_baseline/.

    python scripts/extract_v1.py            # regenerate data/v1_baseline/
    python scripts/extract_v1.py --check    # regenerate into a temp dir and fail if outputs differ

Design rules
- Faithful: every value comes from the workbook. Changes are limited to (a) documented global value maps
  (VALUE_MAPS below) and (b) per-cell text clean-ups, each recorded in normalization_log.csv.
- Derived/suggested values are in columns whose names start with `suggested_` or `derived_`.
- Deterministic: same inputs -> byte-identical CSV/JSON. Timestamps live only in BUILD_INFO.json.
- Provisional stable keys: `v1:<code>` (domain `v1:1`, skill area `v1:1.1`, competency `v1:1.1.1`),
  `v1:1.1.1:T2` (task), `v1:1.1.1:foundational` (proficiency level), `v1:1.1.1:foundational:I3`
  (indicator), `v1:1.1.1:F5` (family assignment), `v1:1.1.1:L2` (ethical lens),
  `v1:1.1.1:sector:ngos` (sector application). These are v1 provenance keys, not v2 IDs.

Only openpyxl and the standard library are required.
"""
from __future__ import annotations

import argparse
import csv
import difflib
import filecmp
import hashlib
import io
import json
import re
import shutil
import sys
import tempfile
from collections import Counter, OrderedDict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import v1_exposure_records  # noqa: E402
import v1_source_text  # noqa: E402

SCRIPT_VERSION = "1.0.0"
ROOT = Path(__file__).resolve().parents[1]
LEGACY = ROOT / "legacy" / "v1" / "data"
OUT = ROOT / "data" / "v1_baseline"

BASELINE = "20260223_SkillsTaxonomy_Core_V3_reconciled__copyB_aaf34c8c.xlsx"
AUX = {
    "d5": "domain5_task13_exposure_table.xlsx",
    "d6": "domain6_task13_exposure_classification.xlsx",
    "pre15a": "PreTask15A_D6_Ac_Reclassification.xlsx",
}
SOURCE_TEXT_GLOB = "20251202_SkillsTaxonomy_Core_V2*.txt"   # in legacy/v1/docs/ (exported by the docs helper)
# Other database versions: hashed for provenance (diffed by scripts/diff_v1_versions.py), not read here.
COMPARED = [
    "skills_taxonomy_workforce_v1.xlsx",
    "skills_taxonomy_workforce_v1_updated.xlsx",
    "20260223_SkillsTaxonomy_Core_V3.xlsx",
    "20260223_SkillsTaxonomy_Core_V3_reconciled__copyA_995a86e4.xlsx",
]

LEVEL_ORDER = {"foundational": 1, "intermediate": 2, "advanced": 3, "expert": 4}
LEVELS = [("FOUNDATIONAL", "foundational", 1), ("INTERMEDIATE", "intermediate", 2),
          ("ADVANCED", "advanced", 3), ("EXPERT", "expert", 4)]
SECTORS = [("Public Agencies", "public_agencies"), ("Private Sector", "private_sector"), ("NGOs", "ngos"),
           ("Consulting", "consulting"), ("Academia", "academia")]
REFORMS = [("Reform: NEPA Timeline", "nepa_timeline"), ("Reform: Digital Permitting", "digital_permitting"),
           ("Reform: CatEx Scope", "catex_scope"), ("Reform: Consultation", "consultation"),
           ("Reform: One Federal Decision", "one_federal_decision"), ("Reform: S2-S3 Boundary", "s2_s3_boundary")]
LENSES = [("L1 Equity & Justice", "L1", "Equity & Justice"),
          ("L2 Cultural Rights & Self-Determination", "L2", "Cultural Rights & Self-Determination"),
          ("L3 Transparency & Accountability", "L3", "Transparency & Accountability"),
          ("L4 Intervention Ethics", "L4", "Intervention Ethics"),
          ("L5 Professional Integrity", "L5", "Professional Integrity")]
TASK_DIMS = [("Auto Type", "auto_type"), ("Tacit", "tacit_density"), ("Career", "career_ladder"),
             ("Vector", "vector"), ("Horizon", "horizon")]
MODAL = [("T13: Auto Type (modal)", "modal_auto_type"), ("T13: Tacit Density (modal)", "modal_tacit_density"),
         ("T13: Career Ladder (modal)", "modal_career_ladder"), ("T13: Vector (modal)", "modal_vector"),
         ("T13: Horizon (modal)", "modal_horizon"), ("T13: CLD Flag (modal)", "modal_cld_flag"),
         ("T13: Modal Contested", "modal_contested")]

# Controlled vocabularies. `documented` = in the v1 Data Dictionary sheet; `observed_undocumented` = present
# in the baseline but not defined anywhere in v1 (kept verbatim, flagged by the audit).
VOCAB = OrderedDict([
    ("craft", {"documented": {"FC": "Fieldcraft", "DC": "Datacraft", "SC": "Statecraft"}}),
    ("stage", {"documented": {"S1": "Planning", "S2": "Environmental Review", "S3": "Permitting",
                              "S4": "Field Operations", "S5": "Monitoring", "S6": "Adaptive Management & Learning"}}),
    ("stability_tier", {"documented": {"Stable": "No reform flags, no Ac/S tasks, not Emergent",
                                       "Transition": "1+ reform flag or 1+ Ac/S task",
                                       "Emergent": "New professional territory"}}),
    ("family_assignment", {"documented": {"Core": "Regularly required across profession",
                                          "Supporting": "Used by some practitioners"}}),
    ("lens_value", {"documented": {"E": "Explicit", "I": "Implicit", "A": "Absent"}}),
    ("auto_type", {"documented": {"A": "Augmentation", "Ac": "Compressive augmentation", "H": "Human-essential",
                                  "S": "Substitution", "V": "Supervision"}}),
    ("tacit_density", {"documented": {"Hi": "High", "Med": "Medium", "Lo": "Low"}}),
    ("career_ladder", {"documented": {"P": "Primary learning (0-5 yr)", "C": "Contributing (3-8 yr)",
                                      "X": "Peripheral/Specialist (8+ yr)"}}),
    ("vector", {"documented": {"AI": "Artificial intelligence", "Ro": "Physical robotics",
                               "Cv": "Convergent (AI + robotics)", "none": "None (H only); v1 wrote an em dash"}}),
    ("horizon", {"documented": {"1": "Deployed now", "2": "3-5 years", "3": "5-10+ years",
                                "N": "Human-essential (no automation trajectory)"}}),
    ("cld_severity", {"documented": {"High": "Primary learning mechanism substantially eliminated",
                                     "Moderate": "Developmental value meaningfully compressed",
                                     "Low": "Administrative overhead automated"},
                      "observed_undocumented": {"Highest": "Used by D6 exposure file for S-at-P full task elimination",
                                                "Ungraded": "No definition in v1"}}),
    ("exposure_provenance", {"documented": {"session_direct": "Session-direct",
                                            "reconstructed_d2": "Reconstructed (D2)",
                                            "reconstructed_d6": "Reconstructed (D6)"},
                             "observed_undocumented": {
                                 "d6_exposure_file_reconciled_2026_03": "D6 exposure file, reconciled March 2026",
                                 "pre_task_15a_d6_ac_retrospective": "D6 Ac retrospective, Task Pre-15A"}}),
    ("mapping_quality", {"documented": {"High": "Near 1:1 SOC alignment", "Medium": "Partial",
                                        "Low": "Fragmented or undercounted"}}),
    ("proficiency_level", {"documented": {"foundational": "Foundational", "intermediate": "Intermediate",
                                          "advanced": "Advanced", "expert": "Expert"}}),
    ("sector", {"documented": {k: lbl for lbl, k in SECTORS}}),
    ("reform_code", {"documented": {k: lbl.replace("Reform: ", "") for lbl, k in REFORMS}}),
])

PROVENANCE_MAP = {
    "Session-direct": "session_direct",
    "Reconstructed (D2)": "reconstructed_d2",
    "Reconstructed (D6)": "reconstructed_d6",
    "D6 exposure file, reconciled March 2026": "d6_exposure_file_reconciled_2026_03",
    "D6 Ac retrospective, Task Pre-15A": "pre_task_15a_d6_ac_retrospective",
}
BOOL_MAP = {"TRUE": True, "FALSE": False, "YES": True, "NO": False, "Y": True, "N": False}

# Global value maps applied silently (documented here and in data/v1_baseline/README.md).
VALUE_MAPS = {
    "booleans": "TRUE/FALSE/True/False/Yes/No/Y/N -> true/false",
    "vector": "em dash '—' -> none",
    "craft_secondary": "'DC-sec' -> 'DC' (suffix dropped; column name carries the role)",
    "skill_area_code": "float 1.1 -> string '1.1' (derived from competency code, never from the float)",
    "exposure_provenance": "label -> snake_case code (raw label kept in exposure_provenance_raw)",
    "whitespace": "leading/trailing whitespace stripped; internal runs of spaces collapsed",
}

# Source-document headings that the v1 import appended to the last field of a block (Academia, usually).
SA_TAIL_RE = re.compile(r"\s+(DOMAIN-LEVEL COMPETENCY FRAMEWORK|DOMAIN \d: .*|SKILL AREAS \(.*\))\s*$", re.S)
# D5 import shifted '*Knowledge*:' / '*Application*:' / '*Indicators*:' lines into Indicator 1-3.
MISPARSED_IND_RE = re.compile(r"^(Knowledge|Application|Indicators)\*?:")


# ----------------------------------------------------------------------------- helpers
def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def code_key(code: str):
    return tuple(int(x) for x in code.split("."))


def clean_text(v):
    if v is None:
        return None
    s = re.sub(r"[ \t]+", " ", str(v)).strip()
    return s or None


def to_bool(v, field, key, log):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        return v
    s = str(v).strip().upper()
    if s not in BOOL_MAP:
        raise ValueError(f"{key} {field}: not a boolean: {v!r}")
    out = BOOL_MAP[s]
    if s in ("YES", "NO"):
        log.append(dict(key=key, field=field, raw=str(v), normalized=str(out).lower(),
                        rule="non-standard boolean spelling Yes/No (TRUE/FALSE or Y/N elsewhere) -> true/false"))
    return out


def code_str(v):
    """Codes stored as numbers in Excel (1, 1.1) -> canonical strings."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_sheet(wb, name, header_row=0):
    rows = list(wb[name].iter_rows(values_only=True))
    hdr = [h.strip() if isinstance(h, str) else h for h in rows[header_row]]
    out = []
    for i, r in enumerate(rows[header_row + 1:], start=header_row + 2):  # 1-based Excel row numbers
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in r):
            continue
        out.append((i, dict(zip(hdr, r))))
    return hdr, out


def csv_value(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (list, tuple)):
        return "|".join(str(x) for x in v)
    return v


def write_csv(path: Path, rows: list[dict], columns: list[str]):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=columns, lineterminator="\n", extrasaction="raise")
    w.writeheader()
    for r in rows:
        w.writerow({c: csv_value(r.get(c)) for c in columns})
    path.write_text(buf.getvalue(), encoding="utf-8")


def split_title_description(title: str):
    """Heuristic split of 'Noun phrase Sentence-case description.' (v1 import merged 19 of these).
    Split before the first capitalised word after position 0 that is not an acronym (>=2 capitals)."""
    toks = title.split(" ")
    for i in range(1, len(toks)):
        t = toks[i]
        if t[:1].isupper() and sum(ch.isupper() for ch in t) < 2:
            head, tail = " ".join(toks[:i]), " ".join(toks[i:])
            if len(tail.split()) >= 5 and tail.rstrip().endswith("."):
                return head, tail
            return None
    return None


def strip_marker(text, key, field, log):
    """Remove trailing source markers (' *', ' #') and a heading that bled into one cell."""
    if text is None:
        return None, None
    marker = None
    m = re.search(r"\s([*#])\s*$", text)
    if m:
        marker = m.group(1)
        new = text[: m.start()].rstrip()
        log.append(dict(key=key, field=field, raw=text, normalized=new,
                        rule=f"trailing source marker '{marker}' moved to source_marker column"))
        text = new
    m = SA_TAIL_RE.search(text)
    if m:
        new = text[: m.start()].rstrip()
        log.append(dict(key=key, field=field, raw=text if len(text) < 600 else text[:600] + " …[truncated in log]",
                        normalized=new, rule="removed source-document heading/narrative that bled into the cell"))
        text = new
    return text, marker


# ----------------------------------------------------------------------------- extract
def extract(out_dir: Path) -> dict:
    wb = openpyxl.load_workbook(LEGACY / BASELINE, read_only=True, data_only=True)
    log: list[dict] = []

    _, comp_rows = read_sheet(wb, "Competencies")
    _, dt_rows = read_sheet(wb, "Dimension Tags")
    pf_hdr, pf_rows = read_sheet(wb, "Profession Family View")
    _, pi_rows = read_sheet(wb, "Proficiency Indicators")
    _, ss_rows = read_sheet(wb, "Structure Summary")
    family_cols = [h for h in pf_hdr if isinstance(h, str) and re.match(r"^F\d+: ", h)]
    families = [(h.split(":")[0], h.split(": ", 1)[1]) for h in family_cols]

    dt = {code_str(r["Competency Code"]): (i, r) for i, r in dt_rows}
    comp = {code_str(r["Competency Code"]): (i, r) for i, r in comp_rows}
    assert set(dt) == set(comp), "Competencies and Dimension Tags disagree on codes"
    codes = sorted(comp, key=code_key)

    # ---------------- structure
    structure = []
    sa_meta = {}
    for i, r in ss_rows:
        dnum = code_str(r["Domain #"])
        if dnum is None or not re.match(r"^\d+$", dnum):
            continue
        sa = code_str(r["Skill Area Code"])
        sa_meta[sa] = dict(row=i, name=clean_text(r["Skill Area Name"]), count=r["Competency Count"],
                           codes=[c.strip() for c in str(r["Competency Codes"]).split(",")],
                           domain_num=dnum, domain_name=clean_text(r["Domain Name"]))
    domains = OrderedDict()
    sas = OrderedDict()
    for c in codes:
        r = comp[c][1]
        d = c.split(".")[0]
        sa = ".".join(c.split(".")[:2])
        domains.setdefault(d, dict(name=clean_text(r["Domain Name"]), codes=[]))["codes"].append(c)
        sas.setdefault(sa, dict(name=clean_text(r["Skill Area Name"]), domain=d, codes=[]))["codes"].append(c)
    for d, m in domains.items():
        structure.append(dict(key=f"v1:{d}", level="domain", domain_num=int(d), domain_name=m["name"],
                              skill_area_code=None, skill_area_name=None, skill_area_name_summary_sheet=None,
                              n_skill_areas=sum(1 for s in sas.values() if s["domain"] == d),
                              n_competencies=len(m["codes"]), competency_codes=m["codes"],
                              summary_sheet_count=None, summary_sheet_consistent=None))
        for sa, s in sas.items():
            if s["domain"] != d:
                continue
            meta = sa_meta.get(sa, {})
            consistent = (meta.get("codes") == s["codes"] and int(meta.get("count") or -1) == len(s["codes"]))
            structure.append(dict(key=f"v1:{sa}", level="skill_area", domain_num=int(d), domain_name=m["name"],
                                  skill_area_code=sa, skill_area_name=s["name"],
                                  skill_area_name_summary_sheet=meta.get("name"),
                                  n_skill_areas=None, n_competencies=len(s["codes"]), competency_codes=s["codes"],
                                  summary_sheet_count=meta.get("count"), summary_sheet_consistent=consistent))

    # ---------------- competencies, tasks, sector applications
    competencies, tasks, sector_apps = [], [], []
    for c in codes:
        row, r = comp[c]
        key = f"v1:{c}"
        title = clean_text(r["Competency Title"])
        desc = clean_text(r["Description"])
        flags = []
        desc, dmark = strip_marker(desc, key, "description", log)
        sug_t = sug_d = None
        if desc is None:
            flags.append("description_missing")
            sp = split_title_description(title)
            if sp:
                flags.append("title_contains_description")
                sug_t, sug_d = sp
        if r["Foundational - Knowledge"] is None:
            flags.append("proficiency_statements_missing")
        sector_markers = set()
        for lbl, sk in SECTORS:
            skey = f"{key}:sector:{sk}"
            text, mk = strip_marker(clean_text(r[lbl]), skey, "text", log)
            if mk:
                sector_markers.add(mk)
            sector_apps.append(dict(key=skey, competency_key=key, code=c, sector=sk, sector_label=lbl,
                                    text=text, source_marker=mk, source_row=row))
        if sector_markers:
            flags.append("sector_text_marker_" + "".join(sorted(sector_markers)).replace("*", "asterisk"))
        if dmark:
            flags.append("description_marker_" + {"#": "hash", "*": "asterisk"}[dmark])
        d = c.split(".")[0]
        sa = ".".join(c.split(".")[:2])
        competencies.append(dict(
            key=key, code=c, domain_key=f"v1:{d}", skill_area_key=f"v1:{sa}", domain_num=int(d),
            domain_name=clean_text(r["Domain Name"]), skill_area_code=sa, skill_area_name=clean_text(r["Skill Area Name"]),
            title=title, description=desc, description_source_marker=dmark,
            suggested_title=sug_t, suggested_description=sug_d, qa_flags=flags,
            source_sheet="Competencies", source_row=row))

        drow, t = dt[c]
        prov_raw = clean_text(t["T13: Data Provenance"])
        for n in (1, 2, 3):
            tk = f"{key}:T{n}"
            rec = dict(key=tk, competency_key=key, code=c, task_n=n, text=clean_text(r[f"Job Task {n}"]),
                       exposure_provenance=PROVENANCE_MAP[prov_raw], exposure_provenance_raw=prov_raw,
                       source_row_text=row, source_row_exposure=drow)
            for col, f in TASK_DIMS:
                v = clean_text(code_str(t[f"T13: T{n} {col}"]))
                if f == "vector" and v == "—":
                    v = "none"
                rec[f] = v
            tasks.append(rec)

    # ---------------- proficiency + indicators
    proficiency, indicators = [], []
    lvl = {a: (b, n) for a, b, n in LEVELS}
    for i, r in sorted(pi_rows, key=lambda x: (code_key(code_str(x[1]["Competency Code"])),
                                               lvl[x[1]["Proficiency Level"].strip()][1])):
        c = code_str(r["Competency Code"])
        lname, lnum = lvl[r["Proficiency Level"].strip()]
        pk = f"v1:{c}:{lname}"
        inds = [(n, clean_text(r[f"Indicator {n}"])) for n in range(1, 7)]
        present = [n for n, v in inds if v]
        gaps = bool(present) and present != list(range(1, max(present) + 1))
        pflags = []
        if clean_text(r["Knowledge"]) is None or clean_text(r["Application"]) is None:
            pflags.append("knowledge_or_application_missing")
        if any(v and MISPARSED_IND_RE.match(v) for _, v in inds):
            pflags.append("indicators_misparsed")
        if len(present) == 6:
            pflags.append("indicators_at_import_cap_6")
        proficiency.append(dict(key=pk, competency_key=f"v1:{c}", code=c, level=lname, level_num=lnum,
                                knowledge=clean_text(r["Knowledge"]), application=clean_text(r["Application"]),
                                n_indicators=len(present), n_indicators_valid=sum(
                                    1 for _, v in inds if v and not MISPARSED_IND_RE.match(v)),
                                indicator_numbering_gap=gaps, qa_flags=pflags,
                                source_sheet="Proficiency Indicators", source_row=i))
        for n, v in inds:
            if v:
                mm = MISPARSED_IND_RE.match(v)
                qa = f"misparsed_{mm.group(1).lower()}_line" if mm else None
                indicators.append(dict(key=f"{pk}:I{n}", proficiency_key=pk, competency_key=f"v1:{c}", code=c,
                                       level=lname, level_num=lnum, indicator_n=n, text=v, qa_flag=qa, source_row=i))

    # ---------------- craft / stage / reform, families, lenses, exposure summary
    csr, fam_rows, lens_rows, exposure = [], [], [], []
    for c in codes:
        row, t = dt[c]
        key = f"v1:{c}"
        sec = clean_text(t["Craft Secondary"])
        rec = dict(key=key, code=c, craft_primary=clean_text(t["Craft Primary"]),
                   craft_secondary=sec.replace("-sec", "") if sec else None,
                   stage_primary_1=clean_text(t["Stage Primary 1"]), stage_primary_2=clean_text(t["Stage Primary 2"]),
                   stage_secondary_1=clean_text(t["Stage Secondary 1"]),
                   stage_secondary_2=clean_text(t["Stage Secondary 2"]),
                   dual_primary_exception=to_bool(t["Dual Primary Exception"], "dual_primary_exception", key, log))
        for col, f in REFORMS:
            rec[f"reform_{f}"] = to_bool(t[col], f"reform_{f}", key, log)
        rec.update(reform_flag_count=t["Reform Flag Count"], stability_tier=clean_text(t["Stability Tier"]),
                   foundational_professional_prep=to_bool(t["Foundational Professional Prep"],
                                                          "foundational_professional_prep", key, log),
                   zero_core=to_bool(t["Zero Core"], "zero_core", key, log),
                   core_family_count=t["Core Family Count"], supporting_family_count=t["Supporting Family Count"],
                   source_sheet="Dimension Tags", source_row=row)
        csr.append(rec)
        for (fcode, fname), col in zip(families, family_cols):
            v = clean_text(t[col])
            if v:
                fam_rows.append(dict(key=f"{key}:{fcode}", competency_key=key, code=c, family_code=fcode,
                                     family_name=fname, assignment=v, source_row=row))
        for col, lcode, lname in LENSES:
            v = clean_text(t[col])
            if v:
                lens_rows.append(dict(key=f"{key}:{lcode}", competency_key=key, code=c, lens_code=lcode,
                                      lens_name=lname, value=v, source_row=row))
        prov_raw = clean_text(t["T13: Data Provenance"])
        ex = dict(key=key, code=c,
                  cld_flag=to_bool(t["T13: CLD Flag"], "cld_flag", key, log),
                  s_trajectory_flag=to_bool(t["T13: S-Trajectory Flag"], "s_trajectory_flag", key, log),
                  cld_severity=clean_text(t["T13: CLD Severity"]))
        for col, f in MODAL:
            v = clean_text(code_str(t[col]))
            if f in ("modal_cld_flag", "modal_contested"):
                v = to_bool(v, f, key, log)
            elif f == "modal_vector" and v == "—":
                v = "none"
            ex[f] = v
        ex.update(exposure_provenance=PROVENANCE_MAP[prov_raw], exposure_provenance_raw=prov_raw,
                  source_sheet="Dimension Tags", source_row=row)
        exposure.append(ex)

    # Profession Family View must mirror Dimension Tags (checked, not extracted twice)
    pf_mismatch = 0
    for i, r in pf_rows:
        c = code_str(r["Competency Code"])
        t = dt[c][1]
        for col in family_cols + ["Craft Primary", "Stability Tier"]:
            if clean_text(r[col]) != clean_text(t[col]):
                pf_mismatch += 1

    # ---------------- labor market crosswalk
    _, cw_rows = read_sheet(wb, "Labor Market Crosswalk", header_row=3)
    crosswalk = []
    for i, r in cw_rows:
        fc = clean_text(r["Family Code"])
        crosswalk.append(dict(key=f"v1:xwalk:{fc}", family_code=fc, family_name=clean_text(r["Family Name"]),
                              primary_soc_code=clean_text(r["Primary SOC Code"]),
                              primary_soc_title=clean_text(r["Primary SOC Title"]),
                              secondary_soc_code=clean_text(r["Secondary SOC Code"]),
                              secondary_soc_title=clean_text(r["Secondary SOC Title"]),
                              mapping_notes=clean_text(r["Mapping Notes"]),
                              mapping_quality=clean_text(r["Mapping Quality"]),
                              soc_vintage_assumed="SOC 2018", source_row=i))

    # ---------------- workforce signals (dashboard sheet -> long table)
    ws_rows = list(wb["Workforce Signals"].iter_rows(values_only=True))
    signals = []
    block = sub = None
    header = None
    for i, r in enumerate(ws_rows, start=1):
        cells = [x for x in r if x is not None and str(x).strip() != ""]
        if not cells:
            header = None
            continue
        if len(cells) == 1 and isinstance(cells[0], str) and cells[0].strip().isupper() and i > 2:
            block, sub, header = clean_text(cells[0]), None, None
            continue
        if i <= 2:
            continue
        if len(cells) == 1 and block and header is None:
            sub = clean_text(cells[0])
            continue
        if header is None:
            header = [clean_text(x) for x in r]
            continue
        label = clean_text(code_str(r[0]))
        for j in range(1, len(r)):
            if r[j] is None or header[j] is None:
                continue
            signals.append(dict(key=f"v1:signal:{len(signals) + 1:03d}", block=block, block_note=sub,
                                row_label=label, row_label_header=header[0], column=header[j],
                                value=clean_text(code_str(r[j])), source_row=i))

    # ---------------- data dictionary
    dd_rows = list(wb["Data Dictionary"].iter_rows(values_only=True))
    ddict = []
    section = "COLUMN DEFINITIONS"
    for i, r in enumerate(dd_rows, start=1):
        cells = [clean_text(code_str(x)) for x in r]
        nonnull = [x for x in cells if x]
        if i <= 3 or not nonnull:
            continue
        if len(nonnull) == 1 and nonnull[0].split("(")[0].strip().isupper() and len(nonnull[0]) > 8:
            section = nonnull[0]
            continue
        is_def_section = section == "COLUMN DEFINITIONS" or section.startswith("TASK 13")
        entry_type = "field" if is_def_section and cells[1] and len(nonnull) >= 3 else "note"
        ddict.append(dict(key=f"v1:dd:{i:03d}", section=section, entry_type=entry_type,
                          field=cells[0], sheet=cells[1] if len(cells) > 1 else None,
                          values_vocabulary=cells[2] if len(cells) > 2 else None,
                          notes=cells[3] if len(cells) > 3 else None, source_row=i))

    # ---------------- auxiliary per-task exposure sources (rationale + cross-check)
    base_task = {t["key"]: t for t in tasks}
    aux_rows = []

    def add_aux(src, sheet, i, code, tn, text, vals, extra):
        tk = f"v1:{code}:{tn}"
        b = base_task.get(tk)
        vec = vals[3]
        vec = "none" if vec == "—" else vec
        vals = (vals[0], vals[1], vals[2], vec, vals[4])
        match = None
        sim = None
        if b:
            match = vals == tuple(b[f] for _, f in TASK_DIMS)
            sim = round(difflib.SequenceMatcher(None, (text or "").lower(), (b["text"] or "").lower()).ratio(), 3)
        aux_rows.append(dict(key=f"{tk}:{src}", task_key=tk, source=src, source_file=AUX[src], source_sheet=sheet,
                             source_row=i, source_task_text=text, task_text_similarity=sim,
                             auto_type=vals[0], tacit_density=vals[1], career_ladder=vals[2], vector=vals[3],
                             horizon=vals[4], matches_baseline=match, **extra))

    w5 = openpyxl.load_workbook(LEGACY / AUX["d5"], read_only=True, data_only=True)
    _, r5 = read_sheet(w5, "Domain5_Task13_Exposure")
    for i, r in r5:
        if not r.get("Comp Code") or not str(r.get("Task #") or "").startswith("T"):
            continue
        add_aux("d5", "Domain5_Task13_Exposure", i, code_str(r["Comp Code"]), r["Task #"].strip(),
                clean_text(r["Task Description"]),
                tuple(clean_text(code_str(r[k])) for k in ("Auto Type", "Tacit Density", "Career Ladder", "Vector",
                                                            "Time Horizon")),
                dict(prev_auto_type=None, cld_flag=bool(clean_text(r["CLD Flag"])), cld_severity=None,
                     delphi_flag=bool(clean_text(r["Delphi Flag"])), rationale=clean_text(r["Notes"])))
    w6 = openpyxl.load_workbook(LEGACY / AUX["d6"], read_only=True, data_only=True)
    _, r6 = read_sheet(w6, "D6 Task Classifications")
    for i, r in r6:
        if not r.get("Task") or not str(r["Task"]).startswith("T"):
            continue
        add_aux("d6", "D6 Task Classifications", i, code_str(r["Code"]), r["Task"].strip(),
                clean_text(r["Task Description"]),
                tuple(clean_text(code_str(r[k])) for k in ("Auto\nType", "Tacit\nDensity", "Career\nLadder", "Vector",
                                                            "Horizon")),
                dict(prev_auto_type=None, cld_flag="CLD" in (r["Notes"] or ""), cld_severity=None,
                     delphi_flag="Delphi" in (r["Notes"] or ""), rationale=clean_text(r["Notes"])))
    wp = openpyxl.load_workbook(LEGACY / AUX["pre15a"], read_only=True, data_only=True)
    _, rp = read_sheet(wp, "Reclassification Table")
    for i, r in rp:
        if not r.get("Code"):
            continue
        add_aux("pre15a", "Reclassification Table", i, code_str(r["Code"]), r["Task"].strip(),
                clean_text(r["Description"]),
                tuple(clean_text(code_str(r[k])) for k in ("New", "Tacit", "Ladder", "Vector", "Horizon")),
                dict(prev_auto_type=clean_text(r["Prev"]), cld_flag=to_bool(r["CLD?"], "cld", "pre15a", []),
                     cld_severity=clean_text(r["CLD Severity"]),
                     delphi_flag=to_bool(r["Delphi?"], "delphi", "pre15a", []), rationale=clean_text(r["Notes"])))
    # per-domain working records in legacy/v1/docs (Task 12/13/14 rationale; D2 confidence codes; D3 inference)
    rec_rows, rec_missing = v1_exposure_records.parse_all(ROOT / "legacy" / "v1" / "docs")
    seen_rec = Counter()
    for x in rec_rows:
        tk = f"v1:{x['code']}:{x['task']}"
        seen_rec[(x["source"], tk)] += 1
        b = base_task.get(tk)
        vals = tuple(x[f] for _, f in TASK_DIMS)
        aux_rows.append(dict(
            key=f"{tk}:{x['source']}" + (f"#{seen_rec[(x['source'], tk)]}" if seen_rec[(x["source"], tk)] > 1 else ""),
            task_key=tk, source=x["source"], source_file=f"legacy/v1/docs/{x['source_file']}", source_sheet=None,
            source_row=x["source_row"], source_task_text=x["source_task_text"],
            task_text_similarity=_sim(x["source_task_text"], b["text"]) if b else None,
            **{f: v for (_, f), v in zip(TASK_DIMS, vals)},
            matches_baseline=(vals == tuple(b[f] for _, f in TASK_DIMS)) if b else None,
            prev_auto_type=None, cld_flag=None, cld_severity=None, delphi_flag="Delphi" in (x["rationale"] or ""),
            rationale=x["rationale"], confidence=x.get("confidence"), source_inferred=x.get("source_inferred")))
    for x in aux_rows:
        if x["source_file"] in AUX.values():
            x["source_file"] = f"legacy/v1/data/{x['source_file']}"
    aux_rows.sort(key=lambda x: (code_key(x["task_key"].split(":")[1]), x["task_key"], x["source"], x["key"]))

    # ---------------- vocabularies
    vocab_rows = []
    for field, spec in VOCAB.items():
        for status in ("documented", "observed_undocumented"):
            for code, label in spec.get(status, {}).items():
                vocab_rows.append(dict(vocabulary=field, code=code, label=label, status=status))
    for fcode, fname in families:
        vocab_rows.append(dict(vocabulary="family", code=fcode, label=fname, status="documented"))
    for _, lcode, lname in LENSES:
        vocab_rows.append(dict(vocabulary="lens", code=lcode, label=lname, status="documented"))

    log.sort(key=lambda x: (x["key"], x["field"], x["rule"]))

    tables = OrderedDict([
        ("structure", (structure, ["key", "level", "domain_num", "domain_name", "skill_area_code", "skill_area_name",
                                   "skill_area_name_summary_sheet", "n_skill_areas", "n_competencies",
                                   "competency_codes", "summary_sheet_count", "summary_sheet_consistent"])),
        ("competencies", (competencies, ["key", "code", "domain_key", "skill_area_key", "domain_num", "domain_name",
                                         "skill_area_code", "skill_area_name", "title", "description",
                                         "description_source_marker", "suggested_title", "suggested_description",
                                         "qa_flags", "source_sheet", "source_row"])),
        ("tasks", (tasks, ["key", "competency_key", "code", "task_n", "text", "auto_type", "tacit_density",
                           "career_ladder", "vector", "horizon", "exposure_provenance", "exposure_provenance_raw",
                           "source_row_text", "source_row_exposure"])),
        ("proficiency", (proficiency, ["key", "competency_key", "code", "level", "level_num", "knowledge",
                                       "application", "n_indicators", "n_indicators_valid", "indicator_numbering_gap",
                                       "qa_flags", "source_sheet", "source_row"])),
        ("indicators", (indicators, ["key", "proficiency_key", "competency_key", "code", "level", "level_num",
                                     "indicator_n", "text", "qa_flag", "source_row"])),
        ("craft_stage_reform", (csr, ["key", "code", "craft_primary", "craft_secondary", "stage_primary_1",
                                      "stage_primary_2", "stage_secondary_1", "stage_secondary_2",
                                      "dual_primary_exception"] + [f"reform_{f}" for _, f in REFORMS] +
                                ["reform_flag_count", "stability_tier", "foundational_professional_prep", "zero_core",
                                 "core_family_count", "supporting_family_count", "source_sheet", "source_row"])),
        ("family_assignments", (fam_rows, ["key", "competency_key", "code", "family_code", "family_name",
                                           "assignment", "source_row"])),
        ("ethical_lenses", (lens_rows, ["key", "competency_key", "code", "lens_code", "lens_name", "value",
                                        "source_row"])),
        ("exposure_summary", (exposure, ["key", "code", "cld_flag", "s_trajectory_flag", "cld_severity"] +
                              [f for _, f in MODAL] + ["exposure_provenance", "exposure_provenance_raw",
                                                       "source_sheet", "source_row"])),
        ("exposure_sources", (aux_rows, ["key", "task_key", "source", "source_file", "source_sheet", "source_row",
                                         "source_task_text", "task_text_similarity", "auto_type", "tacit_density",
                                         "career_ladder", "vector", "horizon", "matches_baseline", "prev_auto_type",
                                         "cld_flag", "cld_severity", "delphi_flag", "confidence", "source_inferred",
                                         "rationale"])),
        ("sector_applications", (sector_apps, ["key", "competency_key", "code", "sector", "sector_label", "text",
                                               "source_marker", "source_row"])),
        ("labor_market_crosswalk", (crosswalk, ["key", "family_code", "family_name", "primary_soc_code",
                                                "primary_soc_title", "secondary_soc_code", "secondary_soc_title",
                                                "mapping_notes", "mapping_quality", "soc_vintage_assumed",
                                                "source_row"])),
        ("workforce_signals", (signals, ["key", "block", "block_note", "row_label", "row_label_header", "column",
                                         "value", "source_row"])),
        ("data_dictionary", (ddict, ["key", "section", "entry_type", "field", "sheet", "values_vocabulary", "notes",
                                     "source_row"])),
        ("vocabularies", (vocab_rows, ["vocabulary", "code", "label", "status"])),
        ("normalization_log", (log, ["key", "field", "raw", "normalized", "rule"])),
    ])

    out_dir.mkdir(parents=True, exist_ok=True)
    for name, (rows, cols) in tables.items():
        write_csv(out_dir / f"{name}.csv", rows, cols)

    # ---------------- combined JSON (nested, one object per competency)
    by = lambda rows, k="competency_key": {r[k]: r for r in rows}  # noqa: E731
    group = {}
    for name in ("tasks", "proficiency", "indicators", "family_assignments", "ethical_lenses", "sector_applications"):
        g = {}
        for r in tables[name][0]:
            g.setdefault(r.get("competency_key"), []).append(r)
        group[name] = g
    ind_by_p = {}
    for r in indicators:
        ind_by_p.setdefault(r["proficiency_key"], []).append({"n": r["indicator_n"], "text": r["text"]})
    csr_by, ex_by = by(csr, "key"), by(exposure, "key")
    comps_json = []
    for cr in competencies:
        k = cr["key"]
        tags = {x: y for x, y in csr_by[k].items() if x not in ("key", "code", "source_sheet", "source_row")}
        exs = {x: y for x, y in ex_by[k].items() if x not in ("key", "code", "source_sheet", "source_row")}
        comps_json.append(OrderedDict(
            **{x: y for x, y in cr.items()},
            tasks=[{x: y for x, y in t.items() if x not in ("competency_key", "code")} for t in group["tasks"].get(k, [])],
            proficiency=[dict(level=p["level"], level_num=p["level_num"], knowledge=p["knowledge"],
                              application=p["application"], indicators=ind_by_p.get(p["key"], []), key=p["key"])
                         for p in group["proficiency"].get(k, [])],
            sector_applications={s["sector"]: s["text"] for s in group["sector_applications"].get(k, [])},
            tags=tags,
            families={f["family_code"]: f["assignment"] for f in group["family_assignments"].get(k, [])},
            ethical_lenses={x["lens_code"]: x["value"] for x in group["ethical_lenses"].get(k, [])},
            exposure=exs))
    combined = OrderedDict(
        dataset="EMST v1 baseline (normalized)", extractor_version=SCRIPT_VERSION,
        baseline_workbook=BASELINE, baseline_sha256=sha256(LEGACY / BASELINE),
        counts=OrderedDict((n, len(t[0])) for n, t in tables.items()),
        structure=structure, competencies=comps_json, labor_market_crosswalk=crosswalk,
        workforce_signals=signals, data_dictionary=ddict, vocabularies=vocab_rows, value_maps=VALUE_MAPS)
    (out_dir / "v1_baseline.json").write_text(json.dumps(combined, ensure_ascii=False, indent=1, default=str) + "\n",
                                              encoding="utf-8")
    xc = source_crosscheck(out_dir, tables)
    return dict(tables=tables, pf_mismatch=pf_mismatch, source=xc)

def _norm_cmp(s):
    if s is None:
        return ""
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    s = s.replace("–", "-").replace("—", "-")
    return re.sub(r"[\s.;,]+$", "", re.sub(r"\s+", " ", s)).strip().lower()


def _sim(a, b):
    return round(difflib.SequenceMatcher(None, _norm_cmp(a), _norm_cmp(b)).ratio(), 3)


def find_source_text():
    docs = ROOT / "legacy" / "v1" / "docs"
    hits = sorted(docs.glob(SOURCE_TEXT_GLOB)) if docs.exists() else []
    return hits[0] if hits else None


def source_crosscheck(out_dir: Path, tables: dict):
    """Compare the database with the source text it was imported from. Writes data/v1_baseline/source_text/."""
    src = find_source_text()
    if src is None:
        print("WARNING: source text not found; skipping source_text/ cross-check", file=sys.stderr)
        return None
    r = v1_source_text.parse(src)
    sd = out_dir / "source_text"
    sd.mkdir(parents=True, exist_ok=True)
    comps = {c["code"]: c for c in tables["competencies"][0]}
    tasks = {(t["code"], t["task_n"]): t for t in tables["tasks"][0]}
    prof = {(p["code"], p["level"]): p for p in tables["proficiency"][0]}
    inds = {(i["code"], i["level"], i["indicator_n"]): i["text"] for i in tables["indicators"][0]}
    inds_by_level = {}
    for (c, l, n), t in inds.items():
        inds_by_level.setdefault((c, l), set()).add(_norm_cmp(t))
    sect = {(x["code"], x["sector_label"]): x["text"] for x in tables["sector_applications"][0]}
    ck = lambda c: code_key(c)  # noqa: E731

    rows_c = []
    for c in sorted(r["competencies"], key=lambda x: ck(x["code"])):
        db = comps.get(c["code"])
        hdr = c["header_text"] or ""
        db_title = db["title"] if db else None
        db_desc = db["description"] if db else None
        merged = bool(db and db_desc and _norm_cmp(hdr).startswith(_norm_cmp(db_title)) and len(hdr) > len(db_title) + 15)
        src_desc = c["description_lines"]
        if not src_desc and merged:
            src_desc = hdr[len(db_title):].strip()
        rows_c.append(dict(key=f"v1:{c['code']}", code=c["code"], source_line=c["source_line"],
                           source_header_text=hdr, source_description=src_desc, db_title=db_title,
                           db_description=db_desc, in_db=db is not None,
                           title_matches=bool(db) and (_norm_cmp(hdr) == _norm_cmp(db_title) or merged),
                           header_has_inline_description=merged or (db is not None and db_desc is None),
                           description_similarity=_sim(src_desc, db_desc) if db_desc else None))
    rows_t = []
    for t in sorted(r["tasks"], key=lambda x: (ck(x["code"]), x["task_n"])):
        db = tasks.get((t["code"], t["task_n"]))
        rows_t.append(dict(key=f"v1:{t['code']}:T{t['task_n']}", code=t["code"], task_n=t["task_n"],
                           source_text=t["text"], db_text=db["text"] if db else None,
                           exact_match=bool(db) and _norm_cmp(t["text"]) == _norm_cmp(db["text"]),
                           similarity=_sim(t["text"], db["text"]) if db else None))
    rows_p, rows_i = [], []
    for p in sorted(r["proficiency"], key=lambda x: (ck(x["code"]), LEVEL_ORDER[x["level"]])):
        db = prof.get((p["code"], p["level"]))
        rows_p.append(dict(key=f"v1:{p['code']}:{p['level']}", code=p["code"], level=p["level"],
                           source_level_label=p["level_label"], source_knowledge=p["knowledge"],
                           source_application=p["application"], db_knowledge=db["knowledge"] if db else None,
                           db_application=db["application"] if db else None,
                           knowledge_match=bool(db) and _norm_cmp(p["knowledge"]) == _norm_cmp(db["knowledge"]),
                           application_match=bool(db) and _norm_cmp(p["application"]) == _norm_cmp(db["application"]),
                           n_indicators_source=p["n_indicators"], n_indicators_db=db["n_indicators"] if db else None))
    for i in sorted(r["indicators"], key=lambda x: (ck(x["code"]), LEVEL_ORDER[x["level"]], x["indicator_n"])):
        dbt = inds.get((i["code"], i["level"], i["indicator_n"]))
        rows_i.append(dict(key=f"v1:{i['code']}:{i['level']}:I{i['indicator_n']}", code=i["code"], level=i["level"],
                           indicator_n=i["indicator_n"], text=i["text"],
                           in_db_same_position=dbt is not None and _norm_cmp(dbt) == _norm_cmp(i["text"]),
                           in_db_any_position=_norm_cmp(i["text"]) in inds_by_level.get((i["code"], i["level"]), set()),
                           db_text_same_position=dbt))
    rows_s = []
    for x in sorted(r["sectors"], key=lambda x: (ck(x["code"]), x["sector_label"])):
        dbt = sect.get((x["code"], x["sector_label"]))
        sk = dict(SECTORS)[x["sector_label"]]
        rows_s.append(dict(key=f"v1:{x['code']}:sector:{sk}", code=x["code"], sector=sk, source_text=x["text"],
                           db_text=dbt, exact_match=_norm_cmp(x["text"]) == _norm_cmp(dbt)))
    rows_sa = [dict(key=f"v1:{x['skill_area_code']}", skill_area_code=x["skill_area_code"], source_name=x["name"],
                    marked_new_in_source=x["marked_new"], source_line=x["source_line"]) for x in r["skill_areas"]]
    out = OrderedDict([
        ("competencies", (rows_c, ["key", "code", "source_line", "source_header_text", "source_description",
                                   "db_title", "db_description", "in_db", "title_matches",
                                   "header_has_inline_description", "description_similarity"])),
        ("tasks", (rows_t, ["key", "code", "task_n", "source_text", "db_text", "exact_match", "similarity"])),
        ("proficiency", (rows_p, ["key", "code", "level", "source_level_label", "source_knowledge",
                                  "source_application", "db_knowledge", "db_application", "knowledge_match",
                                  "application_match", "n_indicators_source", "n_indicators_db"])),
        ("indicators", (rows_i, ["key", "code", "level", "indicator_n", "text", "in_db_same_position",
                                 "in_db_any_position", "db_text_same_position"])),
        ("sector_applications", (rows_s, ["key", "code", "sector", "source_text", "db_text", "exact_match"])),
        ("skill_areas", (rows_sa, ["key", "skill_area_code", "source_name", "marked_new_in_source", "source_line"])),
    ])
    for name, (rows, cols) in out.items():
        write_csv(sd / f"{name}.csv", rows, cols)
    return dict(path=src, tables=out, n_lines=r["n_lines"])


def build_info(out_dir: Path, result: dict) -> dict:
    inputs = [dict(file=f"legacy/v1/data/{BASELINE}", role="baseline", sha256=sha256(LEGACY / BASELINE),
                   bytes=(LEGACY / BASELINE).stat().st_size)]
    for k, f in AUX.items():
        inputs.append(dict(file=f"legacy/v1/data/{f}", role=f"auxiliary:{k}", sha256=sha256(LEGACY / f),
                           bytes=(LEGACY / f).stat().st_size))
    if result.get("source"):
        sp = result["source"]["path"]
        inputs.append(dict(file=str(sp.relative_to(ROOT)), role="source-text (cross-check)", sha256=sha256(sp),
                           bytes=sp.stat().st_size))
    for f in COMPARED:
        inputs.append(dict(file=f"legacy/v1/data/{f}", role="compared-version (not extracted)",
                           sha256=sha256(LEGACY / f), bytes=(LEGACY / f).stat().st_size))
    outputs = []
    for p in sorted(out_dir.glob("*.csv")) + sorted(out_dir.glob("source_text/*.csv")) + [out_dir / "v1_baseline.json"]:
        outputs.append(dict(file=str(p.relative_to(out_dir)), sha256=sha256(p), bytes=p.stat().st_size))
    return OrderedDict(
        extractor="scripts/extract_v1.py", extractor_version=SCRIPT_VERSION,
        generated_utc=datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        python=sys.version.split()[0], openpyxl=openpyxl.__version__,
        inputs=inputs, row_counts=OrderedDict((n, len(t[0])) for n, t in result["tables"].items()),
        source_text_row_counts=(OrderedDict((n, len(t[0])) for n, t in result["source"]["tables"].items())
                                if result.get("source") else None),
        checks=dict(profession_family_view_mismatches_vs_dimension_tags=result["pf_mismatch"]),
        value_maps=VALUE_MAPS, outputs=outputs,
        note="CSV/JSON outputs are deterministic for fixed inputs; only generated_utc varies between runs.")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--check", action="store_true", help="rebuild in a temp dir and compare with --out")
    a = ap.parse_args(argv)
    if a.check:
        tmp = Path(tempfile.mkdtemp())
        try:
            extract(tmp)
            bad = [str(p.relative_to(tmp)) for p in sorted(tmp.rglob("*")) if p.is_file() and
                   (not (a.out / p.relative_to(tmp)).exists() or
                    not filecmp.cmp(p, a.out / p.relative_to(tmp), shallow=False))]
            if bad:
                print("OUT OF DATE:", ", ".join(bad))
                return 1
            print("v1 baseline is up to date")
            return 0
        finally:
            shutil.rmtree(tmp)
    res = extract(a.out)
    info = build_info(a.out, res)
    (a.out / "BUILD_INFO.json").write_text(json.dumps(info, indent=1) + "\n", encoding="utf-8")
    for n, c in info["row_counts"].items():
        print(f"{n:24s} {c:6d}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
